from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def collect_text(shape, parts: list):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)

    if shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    try:
        is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except NotImplementedError:
        is_group = False
    if is_group:
        for child in shape.shapes:
            collect_text(child, parts)


def extract_pptx_slides(file_path: str):
    try:
        presentation = Presentation(file_path)
    except Exception:
        raise ValueError("Could not read this PowerPoint file. Save it as .pptx and try again.")

    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            collect_text(shape, parts)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append("Speaker notes: " + notes.text.strip())

        text = "\n".join(parts).strip()
        if text:
            slides.append({"page_number": index, "text": text})
    return slides