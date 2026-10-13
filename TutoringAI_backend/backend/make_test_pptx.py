from pptx import Presentation

prs = Presentation()
slides = [
("First Normal Form", "A table is in 1NF when every column holds atomic values and there are no repeating groups."),
("Second Normal Form", "A table is in 2NF when it is in 1NF and has no partial dependency on a composite key."),
("Third Normal Form", "A table is in 3NF when it is in 2NF and has no transitive dependency."),
]
for title, body in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
prs.save("../sample_data/test_slides.pptx")
print("saved")