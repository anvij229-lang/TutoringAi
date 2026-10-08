"""Small local API for the Track D study companion MVP."""
from __future__ import annotations

import json
import math
import mimetypes
import os
import random
import re
import uuid
from collections import Counter
from datetime import datetime, timezone
from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

from pypdf import PdfReader
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
WEB = ROOT / "web"
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
INDEX = DATA / "sources.json"
QUIZ_HISTORY = DATA / "quiz_history.json"
MASTERY_FILE = DATA / "mastery.json"
MAX_BYTES = 25 * 1024 * 1024
PORT = int(os.environ.get("PORT", "8000"))
ALLOWED = {".pdf", ".pptx", ".txt", ".vtt", ".srt"}
ACTIVE_QUIZZES = {}


def load_sources():
    if not INDEX.exists():
        return []
    try:
        return json.loads(INDEX.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []


def save_sources(sources):
    DATA.mkdir(exist_ok=True)
    INDEX.write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")


def clean_name(name):
    return Path(name.replace("\\", "/")).name


def split_text(text, size=900, overlap=120):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(text[start:end].strip())
        if end == len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def topic_for(text, filename):
    # A transparent starter tag: use the first short heading-like line, falling back to the filename.
    for line in text.splitlines():
        line = line.strip()
        if 3 <= len(line) <= 72 and (line.isupper() or line.endswith(":")):
            return line.rstrip(":").title()
    stem = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    return stem[:72].title() or "Uncategorized"


STOP_WORDS = {"about", "above", "after", "again", "also", "among", "and", "any", "are", "because", "been", "before", "being", "between", "but", "can", "could", "does", "each", "explain", "find", "from", "give", "have", "here", "how", "into", "just", "more", "most", "much", "please", "some", "such", "than", "that", "their", "them", "then", "there", "these", "they", "this", "those", "through", "what", "when", "where", "which", "while", "with", "would", "your"}


def terms(text):
    return [word for word in re.findall(r"[\w'-]+", text.casefold()) if len(word) >= 3 and word not in STOP_WORDS]


def definition_concept(question):
    match = re.search(r"\b(?:what\s+(?:is|are)|define|definition\s+of|meaning\s+of)\s+(?:an?\s+|the\s+)?([a-z][\w-]*(?:\s+[a-z][\w-]*){0,2})", question.casefold())
    if not match:
        return None
    concept = re.split(r"\b(?:in|for|on|of|from|to|using|within)\b", match.group(1))[0].strip()
    return concept or None


def is_definition_of(text, concept):
    if not concept:
        return False
    # Match glossary forms such as "Array: ..." and "String : ...".
    return bool(re.search(rf"\b{re.escape(concept)}s?\s*:\s*\S", text, re.IGNORECASE))


def locate(unit):
    if unit["location_type"] == "page":
        return f"page {unit['page']}"
    if unit["location_type"] == "slide":
        return f"slide {unit['slide']}"
    if unit["location_type"] == "timestamp":
        return f"{unit['start']}–{unit['end']}"
    return "transcript"


def retrieve(question, limit=3):
    query_terms = set(terms(question))
    if not query_terms:
        return []
    sources = load_sources()
    all_units = [(source, unit) for source in sources for unit in source["units"]]
    document_frequency = Counter()
    tokenized = []
    for source, unit in all_units:
        counts = Counter(terms(unit["text"]))
        tokenized.append((source, unit, counts))
        document_frequency.update(counts.keys())
    average_length = sum(sum(counts.values()) for _, _, counts in tokenized) / max(len(tokenized), 1)
    concept = definition_concept(question)
    exact_definitions = []
    ranked = []
    for source, unit, counts in tokenized:
        overlap = query_terms.intersection(counts)
        is_definition = is_definition_of(unit["text"], concept)
        if not overlap and not is_definition:
            continue
        length = max(sum(counts.values()), 1)
        # BM25-style weighting reduces the advantage of short title and footer slides.
        score = 0.0
        for term in overlap:
            df = document_frequency[term]
            idf = math.log(1 + (len(tokenized) - df + 0.5) / (df + 0.5))
            tf = counts[term]
            score += idf * (tf * 2.2) / (tf + 1.2 * (0.25 + 0.75 * length / max(average_length, 1)))
        if concept and len(counts) < 6 and not is_definition:
            continue
        if is_definition:
            score += 100.0
            exact_definitions.append((score, source, unit, sorted(overlap)))
        ranked.append((score, source, unit, sorted(overlap)))
    if exact_definitions:
        exact_definitions.sort(key=lambda item: item[0], reverse=True)
        return exact_definitions[:limit]
    ranked.sort(key=lambda item: item[0], reverse=True)
    return [item for item in ranked if item[0] > 0][:limit]


def answer_question(question):
    hits = retrieve(question)
    if not hits or hits[0][0] < 0.12:
        return {"supported": False, "answer": "I couldn't find support for that in your uploaded course material. Try asking about a term or topic that appears in your sources.", "citations": []}
    citations = []
    for score, source, unit, matched in hits:
        citations.append({"source_id": source["id"], "source": source["name"], "location": locate(unit),
                          "topic": unit["topic"], "excerpt": unit["text"], "matched_terms": matched,
                          "relevance": round(score, 3)})
    return {"supported": True, "answer": "I found these passages in your uploaded material. Check the slide citation to see the original context.",
            "citations": citations}


def read_json(path, fallback):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return fallback


def extract_definitions():
    definitions = []
    next_label = re.compile(r"(?<=[.!?;])\s+(?=[A-Z][A-Za-z0-9 /-]{1,45}:\s)")
    label_pattern = re.compile(r"^([A-Z][A-Za-z0-9 /-]{1,45}?)\s*:\s*(.+)$", re.S)
    for source in load_sources():
        for unit in source["units"]:
            for part in next_label.split(unit["text"].strip()):
                match = label_pattern.match(part.strip())
                if not match:
                    continue
                label = re.sub(r"\s+", " ", match.group(1)).strip(" -")
                definition = re.split(r"\s*©|\s*\\\\n", match.group(2))[0].strip(" .;\n")
                if len(terms(label)) < 1 or len(terms(definition)) < 4:
                    continue
                definitions.append({"term": label, "definition": definition, "topic": unit["topic"],
                                    "source_id": source["id"], "source": source["name"], "location": locate(unit)})
    # Avoid duplicate glossary entries from repeated slide text or overlapping chunks.
    unique = {}
    for item in definitions:
        unique.setdefault((item["term"].casefold(), item["definition"].casefold()), item)
    return list(unique.values())


def topic_matches(item, topic):
    wanted = set(terms(topic))
    if not wanted:
        return True
    combined = set(terms(item["term"] + " " + item["definition"] + " " + item["topic"]))
    return bool(wanted.intersection(combined))


def create_quiz(topic="", difficulty="medium", count=5):
    definitions = [item for item in extract_definitions() if topic_matches(item, topic)]
    if not definitions:
        return None
    history = read_json(QUIZ_HISTORY, {"asked": [], "attempts": []})
    already_asked = set(history.get("asked", []))
    fresh = [item for item in definitions if f'{item["term"].casefold()}|{item["definition"].casefold()}' not in already_asked]
    if len(fresh) < count:
        fresh = fresh + [item for item in definitions if item not in fresh]
    random.shuffle(fresh)
    selected = fresh[:count]
    quiz_id = str(uuid.uuid4())
    question_items = []
    for index, item in enumerate(selected):
        short_answer = difficulty == "hard" and index % 2 == 1 or difficulty == "medium" and index % 4 == 3
        question_id = str(uuid.uuid4())
        if short_answer:
            prompt = f'In your own words, define “{item["term"]}”.'
            choices = []
            correct_answer = item["definition"]
        else:
            distractors = [other["term"] for other in definitions if other["term"].casefold() != item["term"].casefold()]
            random.shuffle(distractors)
            choices = list(dict.fromkeys([item["term"], *distractors[:3]]))
            random.shuffle(choices)
            prompt = f'Which term matches this description? “{item["definition"]}”'
            correct_answer = item["term"]
        question_items.append({"id": question_id, "type": "short_answer" if short_answer else "mcq",
                               "prompt": prompt, "choices": choices, "topic": item["term"],
                               "source_id": item["source_id"], "source": item["source"],
                               "location": item["location"], "excerpt": item["definition"],
                               "correct_answer": correct_answer})
        already_asked.add(f'{item["term"].casefold()}|{item["definition"].casefold()}')
    ACTIVE_QUIZZES[quiz_id] = {"questions": question_items, "difficulty": difficulty, "topic": topic}
    history["asked"] = list(already_asked)
    QUIZ_HISTORY.parent.mkdir(exist_ok=True)
    QUIZ_HISTORY.write_text(json.dumps(history, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"quiz_id": quiz_id, "difficulty": difficulty, "topic": topic,
            "questions": [{key: value for key, value in item.items() if key not in {"correct_answer", "excerpt"}}
                          for item in question_items]}


def submit_quiz(quiz_id, answers):
    quiz = ACTIVE_QUIZZES.pop(quiz_id, None)
    if quiz is None:
        return None
    mastery = read_json(MASTERY_FILE, {})
    results = []
    correct_count = 0
    for item in quiz["questions"]:
        response = answers.get(item["id"], {})
        correct_answer = item["correct_answer"]
        if item["type"] == "mcq":
            user_answer = str(response.get("answer", "")).strip()
            is_correct = user_answer.casefold() == correct_answer.casefold()
        else:
            user_answer = str(response.get("answer", "")).strip()
            expected = set(terms(correct_answer))
            supplied = set(terms(user_answer))
            overlap = len(expected.intersection(supplied)) / max(min(len(expected), 8), 1)
            is_correct = len(supplied) >= 2 and overlap >= 0.3
        if is_correct:
            correct_count += 1
        previous = float(mastery.get(item["topic"], 0.5))
        mastery[item["topic"]] = round(previous * 0.7 + (1.0 if is_correct else 0.0) * 0.3, 3)
        results.append({"id": item["id"], "type": item["type"], "prompt": item["prompt"],
                        "user_answer": user_answer, "correct": is_correct,
                        "correct_answer": correct_answer, "explanation": item["excerpt"],
                        "topic": item["topic"], "source_id": item["source_id"],
                        "source": item["source"], "location": item["location"]})
    MASTERY_FILE.parent.mkdir(exist_ok=True)
    MASTERY_FILE.write_text(json.dumps(mastery, ensure_ascii=False, indent=2), encoding="utf-8")
    history = read_json(QUIZ_HISTORY, {"asked": [], "attempts": []})
    attempt = {"submitted_at": datetime.now(timezone.utc).isoformat(), "score": correct_count,
               "total": len(results), "topic": quiz["topic"], "difficulty": quiz["difficulty"]}
    history.setdefault("attempts", []).insert(0, attempt)
    QUIZ_HISTORY.write_text(json.dumps(history, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"score": correct_count, "total": len(results), "results": results,
            "weak_topics": [key for key, value in sorted(mastery.items(), key=lambda pair: pair[1])[:3]],
            "mastery": mastery, "attempt": attempt}


def make_units(ext, raw, filename):
    units = []
    if ext == ".pdf":
        reader = PdfReader(raw)
        for page_no, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            for piece in split_text(text):
                units.append({"location_type": "page", "page": page_no, "text": piece,
                              "topic": topic_for(text, filename)})
    elif ext == ".pptx":
        presentation = Presentation(raw)
        for slide_no, slide in enumerate(presentation.slides, 1):
            lines = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            text = "\n".join(lines)
            for piece in split_text(text):
                units.append({"location_type": "slide", "slide": slide_no, "text": piece,
                              "topic": topic_for(text, filename)})
    else:
        text = raw.read().decode("utf-8-sig", errors="replace")
        # Preserve cue times for VTT/SRT; plain transcripts stay as untimed text units.
        cue_pattern = re.compile(r"(?m)^(\d{2}:\d{2}:\d{2}[.,]\d{3})\s+-->\s+(\d{2}:\d{2}:\d{2}[.,]\d{3}).*?\n(.*?)(?=\n\s*\n|\Z)", re.S)
        cues = list(cue_pattern.finditer(text))
        if cues:
            for cue in cues:
                cue_text = re.sub(r"<[^>]+>", "", cue.group(3)).strip()
                for piece in split_text(cue_text):
                    units.append({"location_type": "timestamp", "start": cue.group(1), "end": cue.group(2),
                                  "text": piece, "topic": topic_for(cue_text, filename)})
        else:
            for piece in split_text(text):
                units.append({"location_type": "transcript", "text": piece,
                              "topic": topic_for(text, filename)})
    return units


class Handler(BaseHTTPRequestHandler):
    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
        self.end_headers()

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/health":
            return self.send_json(200, {"ok": True, "service": "track-d-mvp"})
        if path == "/api/quiz/topics":
            topics = sorted({item["term"] for item in extract_definitions()}, key=str.casefold)
            return self.send_json(200, {"topics": topics})
        if path == "/api/dashboard":
            mastery = read_json(MASTERY_FILE, {})
            history = read_json(QUIZ_HISTORY, {"attempts": []})
            topic_names = sorted({item["term"] for item in extract_definitions()}, key=str.casefold)
            topics = [{"topic": name, "mastery": float(mastery.get(name, 0.5))} for name in topic_names]
            topics.sort(key=lambda item: item["mastery"])
            next_topic = topics[0]["topic"] if topics else None
            return self.send_json(200, {"topics": topics, "attempts": history.get("attempts", [])[:10],
                                        "next_topic": next_topic, "attempt_count": len(history.get("attempts", []))})
        if path == "/api/sources":
            sources = load_sources()
            return self.send_json(200, {"sources": sources, "unit_count": sum(len(s["units"]) for s in sources)})
        file_match = re.fullmatch(r"/api/sources/([0-9a-f-]+)/file", path)
        if file_match:
            source = next((s for s in load_sources() if s["id"] == file_match.group(1)), None)
            if source is None:
                return self.send_json(404, {"error": "Source not found"})
            extension = Path(source["name"]).suffix.lower()
            stored = UPLOADS / f"{source['id']}{extension}"
            if not stored.is_file():
                return self.send_json(404, {"error": "Source file is missing"})
            body = stored.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", mimetypes.guess_type(source["name"])[0] or "application/octet-stream")
            disposition = "inline" if extension == ".pdf" else "attachment"
            self.send_header("Content-Disposition", f'{disposition}; filename="{clean_name(source["name"])}"')
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/" or path.startswith("/web/"):
            target = WEB / ("index.html" if path == "/" else path.removeprefix("/web/"))
            if not target.is_file() or WEB.resolve() not in target.resolve().parents:
                self.send_error(404)
                return
            body = target.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", mimetypes.guess_type(str(target))[0] or "application/octet-stream")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        self.send_error(404)

    def do_POST(self):
        route = urlparse(self.path).path
        quiz_submit = re.fullmatch(r"/api/quizzes/([0-9a-f-]+)/submit", route)
        if route == "/api/quizzes" or quiz_submit:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 32_000:
                return self.send_json(413, {"error": "Quiz request is empty or too large."})
            try:
                payload = json.loads(self.rfile.read(length))
            except (json.JSONDecodeError, UnicodeDecodeError):
                return self.send_json(400, {"error": "Send a JSON object."})
            if route == "/api/quizzes":
                count = min(max(int(payload.get("count", 5)), 1), 10)
                difficulty = str(payload.get("difficulty", "medium")).lower()
                if difficulty not in {"easy", "medium", "hard"}:
                    return self.send_json(400, {"error": "Difficulty must be easy, medium, or hard."})
                quiz = create_quiz(str(payload.get("topic", "")), difficulty, count)
                if quiz is None:
                    return self.send_json(422, {"error": "No glossary questions found for that topic. Try All topics or another topic."})
                return self.send_json(201, quiz)
            result = submit_quiz(quiz_submit.group(1), payload.get("answers", {}))
            if result is None:
                return self.send_json(404, {"error": "Quiz expired. Generate a new quiz and try again."})
            return self.send_json(200, result)
        if route == "/api/chat":
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 32_000:
                return self.send_json(413, {"error": "Question must be between 1 byte and 32 KB."})
            try:
                payload = json.loads(self.rfile.read(length))
                question = str(payload.get("question", "")).strip()
            except (json.JSONDecodeError, UnicodeDecodeError):
                return self.send_json(400, {"error": "Send a JSON object with a question."})
            if not question:
                return self.send_json(400, {"error": "Enter a question first."})
            return self.send_json(200, answer_question(question))
        if route != "/api/sources":
            return self.send_json(404, {"error": "Not found"})
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0 or length > MAX_BYTES:
            return self.send_json(413, {"error": "Upload must be between 1 byte and 25 MB."})
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            return self.send_json(400, {"error": "Upload a file using multipart form data."})
        message = BytesParser(policy=default).parsebytes(
            b"MIME-Version: 1.0\r\nContent-Type: " + content_type.encode() + b"\r\n\r\n" + self.rfile.read(length)
        )
        part = next((p for p in message.iter_parts() if p.get_filename()), None)
        if part is None:
            return self.send_json(400, {"error": "Choose a file to upload."})
        filename = clean_name(part.get_filename())
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED:
            return self.send_json(415, {"error": "Supported files: PDF, PPTX, TXT, VTT, and SRT transcripts."})
        blob = part.get_payload(decode=True) or b""
        UPLOADS.mkdir(parents=True, exist_ok=True)
        source_id = str(uuid.uuid4())
        stored = UPLOADS / f"{source_id}{ext}"
        stored.write_bytes(blob)
        try:
            units = make_units(ext, stored, filename)
        except Exception as exc:
            stored.unlink(missing_ok=True)
            return self.send_json(422, {"error": f"Could not extract this file: {exc}"})
        source = {"id": source_id, "name": filename, "type": ext[1:], "uploaded_at": datetime.now(timezone.utc).isoformat(),
                  "unit_count": len(units), "units": units}
        sources = load_sources()
        sources.insert(0, source)
        save_sources(sources)
        return self.send_json(201, {"source": source})

    def do_DELETE(self):
        source_id = urlparse(self.path).path.removeprefix("/api/sources/")
        sources = load_sources()
        source = next((s for s in sources if s["id"] == source_id), None)
        if source is None:
            return self.send_json(404, {"error": "Source not found"})
        for file in UPLOADS.glob(source_id + ".*"):
            file.unlink(missing_ok=True)
        save_sources([s for s in sources if s["id"] != source_id])
        return self.send_json(200, {"deleted": source_id})

    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")


if __name__ == "__main__":
    DATA.mkdir(exist_ok=True)
    print(f"Track D MVP running at http://127.0.0.1:{PORT}")
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()
